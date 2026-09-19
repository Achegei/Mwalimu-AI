from io import BytesIO

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter

from app.services.document_ingestion import (
    ExtractedPage,
    chunk_extracted_pages,
    extract_document_text,
    normalize_text,
)


def test_normalize_text_collapses_whitespace():
    text = "  Mwalimu   AI \n\n helps\tstudents.  "

    result = normalize_text(text)

    assert result == "Mwalimu AI helps students."


def test_extracts_utf8_plain_text():
    content = (
        "Photosynthesis is the process by which "
        "green plants make food."
    ).encode("utf-8")

    pages = extract_document_text(
        content,
        mime_type="text/plain",
        filename="biology-notes.txt",
    )

    assert len(pages) == 1
    assert pages[0].page_number is None
    assert "Photosynthesis" in pages[0].text


def test_extracts_docx_text():
    buffer = BytesIO()

    document = DocxDocument()
    document.add_heading("Form 1 Biology", level=1)
    document.add_paragraph(
        "Cells are the basic units of life."
    )
    document.save(buffer)

    pages = extract_document_text(
        buffer.getvalue(),
        mime_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        filename="biology.docx",
    )

    assert len(pages) == 1
    assert "Form 1 Biology" in pages[0].text
    assert "Cells are the basic units of life." in pages[0].text


def test_pdf_without_extractable_text_is_rejected():
    buffer = BytesIO()

    writer = PdfWriter()
    writer.add_blank_page(
        width=612,
        height=792,
    )
    writer.write(buffer)

    with pytest.raises(
        ValueError,
        match="extractable text",
    ):
        extract_document_text(
            buffer.getvalue(),
            mime_type="application/pdf",
            filename="scanned-paper.pdf",
        )


def test_unsupported_document_format_is_rejected():
    with pytest.raises(
        ValueError,
        match="Unsupported document format",
    ):
        extract_document_text(
            b"binary-data",
            mime_type="application/octet-stream",
            filename="archive.bin",
        )


def test_chunking_keeps_small_page_as_one_chunk():
    pages = [
        ExtractedPage(
            page_number=1,
            text="A short paragraph about algebra.",
        )
    ]

    chunks = chunk_extracted_pages(
        pages,
        chunk_size=500,
        overlap=50,
    )

    assert len(chunks) == 1
    assert chunks[0].chunk_index == 0
    assert chunks[0].page_number == 1
    assert chunks[0].content == (
        "A short paragraph about algebra."
    )
    assert chunks[0].character_count == len(
        chunks[0].content
    )


def test_chunking_splits_large_text_with_overlap():
    text = " ".join(
        f"word{i}"
        for i in range(300)
    )

    pages = [
        ExtractedPage(
            page_number=4,
            text=text,
        )
    ]

    chunks = chunk_extracted_pages(
        pages,
        chunk_size=300,
        overlap=50,
    )

    assert len(chunks) > 1

    assert [
        chunk.chunk_index
        for chunk in chunks
    ] == list(range(len(chunks)))

    assert all(
        chunk.page_number == 4
        for chunk in chunks
    )

    assert all(
        chunk.character_count == len(chunk.content)
        for chunk in chunks
    )


def test_chunking_preserves_page_boundaries():
    pages = [
        ExtractedPage(
            page_number=1,
            text="Page one content.",
        ),
        ExtractedPage(
            page_number=2,
            text="Page two content.",
        ),
    ]

    chunks = chunk_extracted_pages(
        pages,
        chunk_size=500,
        overlap=50,
    )

    assert len(chunks) == 2
    assert chunks[0].page_number == 1
    assert chunks[1].page_number == 2


@pytest.mark.parametrize(
    ("chunk_size", "overlap"),
    [
        (0, 0),
        (-1, 0),
        (100, -1),
        (100, 100),
        (100, 101),
    ],
)
def test_chunking_rejects_invalid_configuration(
    chunk_size,
    overlap,
):
    with pytest.raises(ValueError):
        chunk_extracted_pages(
            [
                ExtractedPage(
                    page_number=1,
                    text="Some content.",
                )
            ],
            chunk_size=chunk_size,
            overlap=overlap,
        )
