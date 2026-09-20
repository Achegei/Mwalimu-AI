import csv
import io

from fastapi import (
    APIRouter,
    Depends,
    File,
    Form,
    HTTPException,
    UploadFile,
    status,
)
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.core.dependencies import require_roles
from app.models.enums import DocumentType, UserRole
from app.models.user import User
from app.schemas.content import (
    SubjectSummary,
    SubjectWithTopics,
)
from app.schemas.admin import (
    AdminBulkImportResponse,
    AdminClassroomCreate,
    AdminClassroomSummary,
    AdminDocumentSummary,
    AdminEnrollmentStudent,
    AdminTeachingAssignmentCreate,
    AdminTeachingAssignmentSummary,
    AdminUserCreate,
    AdminUserSummary,
)
from app.services.admin_documents import (
    create_school_document,
    get_school_documents,
    validate_document_scope,
)
from app.services.content import (
    get_active_subjects,
    get_active_topics_for_subject,
)
from app.services.teaching_assignments import (
    create_teaching_assignment,
    deactivate_teaching_assignment,
    get_school_teaching_assignments,
)
from app.services.document_processing import process_document
from app.services.document_storage import (
    build_document_storage_key,
    delete_document_file,
    save_document_bytes,
)
from app.services.admin_users import (
    bulk_import_school_students,
    create_school_classroom,
    create_school_user,
    enroll_school_student,
    get_school_classroom_students,
    get_school_classrooms,
    get_school_users,
    remove_school_student_enrollment,
)


router = APIRouter(
    prefix="/admin",
    tags=["Admin"],
)


@router.get(
    "/subjects",
    response_model=list[SubjectSummary],
)
async def list_admin_subjects(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(
        require_roles(UserRole.ADMIN)
    ),
) -> list[SubjectSummary]:
    return await get_active_subjects(
        db=db,
        school_id=current_user.school_id,
    )


@router.get(
    "/subjects/{subject_id}/topics",
    response_model=SubjectWithTopics,
)
async def list_admin_subject_topics(
    subject_id: int,
    form_level: int,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(
        require_roles(UserRole.ADMIN)
    ),
) -> SubjectWithTopics:
    if form_level < 1 or form_level > 4:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Form level must be between 1 and 4.",
        )

    subjects = await get_active_subjects(
        db=db,
        school_id=current_user.school_id,
    )

    subject = next(
        (
            item
            for item in subjects
            if item.id == subject_id
        ),
        None,
    )

    if subject is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Subject not found.",
        )

    topics = await get_active_topics_for_subject(
        db=db,
        subject_id=subject.id,
        school_id=current_user.school_id,
        form_level=form_level,
    )

    return SubjectWithTopics(
        id=subject.id,
        name=subject.name,
        slug=subject.slug,
        description=subject.description,
        topics=topics,
    )


@router.get(
    "/teaching-assignments",
    response_model=list[AdminTeachingAssignmentSummary],
)
async def list_teaching_assignments(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(
        require_roles(UserRole.ADMIN)
    ),
) -> list[AdminTeachingAssignmentSummary]:
    assignments = await get_school_teaching_assignments(
        db=db,
        school_id=current_user.school_id,
    )

    return assignments


@router.post(
    "/teaching-assignments",
    response_model=AdminTeachingAssignmentSummary,
    status_code=status.HTTP_201_CREATED,
)
async def create_admin_teaching_assignment(
    payload: AdminTeachingAssignmentCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(
        require_roles(UserRole.ADMIN)
    ),
) -> AdminTeachingAssignmentSummary:
    try:
        assignment = await create_teaching_assignment(
            db=db,
            school_id=current_user.school_id,
            teacher_id=payload.teacher_id,
            subject_id=payload.subject_id,
            classroom_id=payload.classroom_id,
        )
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(exc),
        ) from exc

    return assignment


@router.delete(
    "/teaching-assignments/{assignment_id}",
    response_model=AdminTeachingAssignmentSummary,
)
async def deactivate_admin_teaching_assignment(
    assignment_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(
        require_roles(UserRole.ADMIN)
    ),
) -> AdminTeachingAssignmentSummary:
    try:
        assignment = await deactivate_teaching_assignment(
            db=db,
            school_id=current_user.school_id,
            assignment_id=assignment_id,
        )
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(exc),
        ) from exc

    return assignment


@router.get(
    "/documents",
    response_model=list[AdminDocumentSummary],
)
async def list_school_documents(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(
        require_roles(UserRole.ADMIN)
    ),
) -> list[AdminDocumentSummary]:
    documents = await get_school_documents(
        db=db,
        school_id=current_user.school_id,
    )

    return documents


@router.post(
    "/documents",
    response_model=AdminDocumentSummary,
    status_code=status.HTTP_201_CREATED,
)
async def upload_school_document(
    title: str = Form(...),
    document_type: DocumentType = Form(...),
    subject_id: int = Form(...),
    form_level: int = Form(...),
    topic_id: int | None = Form(None),
    academic_year: int | None = Form(None),
    exam_year: int | None = Form(None),
    paper_number: str | None = Form(None),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(
        require_roles(UserRole.ADMIN)
    ),
) -> AdminDocumentSummary:
    clean_title = title.strip()

    if not clean_title:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Document title is required.",
        )

    if form_level < 1 or form_level > 4:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Form level must be between 1 and 4.",
        )

    original_filename = (
        file.filename or ""
    ).strip()

    if not original_filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Document filename is required.",
        )

    content = await file.read()

    if not content:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Document file cannot be empty.",
        )

    mime_type = (
        file.content_type
        or "application/octet-stream"
    )

    storage_key = build_document_storage_key(
        school_id=current_user.school_id,
        original_filename=original_filename,
    )

    try:
        # Validate tenant ownership before writing the file.
        await validate_document_scope(
            db=db,
            school_id=current_user.school_id,
            subject_id=subject_id,
            topic_id=topic_id,
            form_level=form_level,
        )

        save_document_bytes(
            storage_key=storage_key,
            content=content,
        )

        document = await create_school_document(
            db=db,
            school_id=current_user.school_id,
            uploaded_by_id=current_user.id,
            subject_id=subject_id,
            topic_id=topic_id,
            title=clean_title,
            document_type=document_type,
            form_level=form_level,
            academic_year=academic_year,
            exam_year=exam_year,
            paper_number=(
                paper_number.strip()
                if paper_number
                else None
            ),
            original_filename=original_filename,
            storage_key=storage_key,
            mime_type=mime_type,
            file_size=len(content),
        )
    except ValueError as exc:
        delete_document_file(storage_key)

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(exc),
        )
    except Exception:
        delete_document_file(storage_key)
        raise

    try:
        await process_document(
            db,
            document.id,
        )
    except (ValueError, OSError):
        # process_document records the failure state on the
        # document. Keep the source file so processing can be
        # retried later.
        pass

    await db.refresh(document)

    return document


@router.get(
    "/users",
    response_model=list[AdminUserSummary],
)
async def list_school_users(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(
        require_roles(UserRole.ADMIN)
    ),
) -> list[AdminUserSummary]:
    users = await get_school_users(
        db=db,
        school_id=current_user.school_id,
    )

    return users


@router.post(
    "/users",
    response_model=AdminUserSummary,
    status_code=status.HTTP_201_CREATED,
)
async def create_user(
    payload: AdminUserCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(
        require_roles(UserRole.ADMIN)
    ),
) -> AdminUserSummary:
    try:
        user = await create_school_user(
            db=db,
            school_id=current_user.school_id,
            login_id=payload.login_id,
            full_name=payload.full_name,
            role=payload.role,
            password=payload.password,
        )
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=str(exc),
        )

    return user


@router.get(
    "/classrooms",
    response_model=list[AdminClassroomSummary],
)
async def list_school_classrooms(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(
        require_roles(UserRole.ADMIN)
    ),
) -> list[AdminClassroomSummary]:
    classrooms = await get_school_classrooms(
        db=db,
        school_id=current_user.school_id,
    )

    return classrooms


@router.post(
    "/classrooms",
    response_model=AdminClassroomSummary,
    status_code=status.HTTP_201_CREATED,
)
async def create_classroom(
    payload: AdminClassroomCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(
        require_roles(UserRole.ADMIN)
    ),
) -> AdminClassroomSummary:
    try:
        classroom = await create_school_classroom(
            db=db,
            school_id=current_user.school_id,
            name=payload.name,
            form_level=payload.form_level,
            academic_year=payload.academic_year,
        )
    except ValueError as exc:
        detail = str(exc)

        if "already exists" in detail:
            status_code = status.HTTP_409_CONFLICT
        else:
            status_code = status.HTTP_400_BAD_REQUEST

        raise HTTPException(
            status_code=status_code,
            detail=detail,
        )

    return classroom


@router.get(
    "/classrooms/{classroom_id}/students",
    response_model=list[AdminEnrollmentStudent],
)
async def list_classroom_students(
    classroom_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(
        require_roles(UserRole.ADMIN)
    ),
) -> list[AdminEnrollmentStudent]:
    try:
        students = await get_school_classroom_students(
            db=db,
            school_id=current_user.school_id,
            classroom_id=classroom_id,
        )
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=str(exc),
        )

    return [
        AdminEnrollmentStudent(**student)
        for student in students
    ]


async def _parse_student_import(
    file: UploadFile,
) -> list[dict[str, str]]:
    filename = (file.filename or "").lower()
    content = await file.read()

    required_columns = {
        "full_name",
        "login_id",
        "password",
    }

    if filename.endswith(".csv"):
        try:
            decoded = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError(
                "CSV file must use UTF-8 encoding."
            ) from exc

        reader = csv.DictReader(
            io.StringIO(decoded)
        )

        fieldnames = {
            str(name).strip()
            for name in (reader.fieldnames or [])
            if name is not None
        }

        if not required_columns.issubset(
            fieldnames
        ):
            raise ValueError(
                "Import file must contain full_name, "
                "login_id, and password columns."
            )

        return [
            {
                str(key).strip(): (
                    "" if value is None else str(value)
                )
                for key, value in row.items()
                if key is not None
            }
            for row in reader
        ]

    if filename.endswith(".xlsx"):
        try:
            workbook = load_workbook(
                io.BytesIO(content),
                read_only=True,
                data_only=True,
            )
        except Exception as exc:
            raise ValueError(
                "Unable to read Excel file."
            ) from exc

        try:
            worksheet = workbook.active
            values = worksheet.iter_rows(
                values_only=True
            )

            header_row = next(values, None)

            if header_row is None:
                raise ValueError(
                    "Import file is empty."
                )

            headers = [
                str(value).strip()
                if value is not None
                else ""
                for value in header_row
            ]

            if not required_columns.issubset(
                set(headers)
            ):
                raise ValueError(
                    "Import file must contain full_name, "
                    "login_id, and password columns."
                )

            rows: list[dict[str, str]] = []

            for values_row in values:
                row = {
                    header: (
                        ""
                        if value is None
                        else str(value)
                    )
                    for header, value in zip(
                        headers,
                        values_row,
                        strict=False,
                    )
                    if header
                }

                rows.append(row)

            return rows
        finally:
            workbook.close()

    raise ValueError(
        "Unsupported file type. Upload a CSV or XLSX file."
    )


@router.post(
    "/classrooms/{classroom_id}/students/import",
    response_model=AdminBulkImportResponse,
)
async def import_classroom_students(
    classroom_id: int,
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(
        require_roles(UserRole.ADMIN)
    ),
) -> AdminBulkImportResponse:
    try:
        rows = await _parse_student_import(file)

        result = await bulk_import_school_students(
            db=db,
            school_id=current_user.school_id,
            classroom_id=classroom_id,
            rows=rows,
        )

    except ValueError as exc:
        detail = str(exc)

        if "Classroom not found" in detail:
            status_code = status.HTTP_404_NOT_FOUND
        else:
            status_code = status.HTTP_400_BAD_REQUEST

        raise HTTPException(
            status_code=status_code,
            detail=detail,
        )

    return AdminBulkImportResponse(**result)


@router.post(
    "/classrooms/{classroom_id}/students/{student_id}",
    response_model=AdminEnrollmentStudent,
    status_code=status.HTTP_201_CREATED,
)
async def enroll_student(
    classroom_id: int,
    student_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(
        require_roles(UserRole.ADMIN)
    ),
) -> AdminEnrollmentStudent:
    try:
        student = await enroll_school_student(
            db=db,
            school_id=current_user.school_id,
            classroom_id=classroom_id,
            student_id=student_id,
        )
    except ValueError as exc:
        detail = str(exc)

        if "already enrolled" in detail:
            status_code = status.HTTP_409_CONFLICT
        else:
            status_code = status.HTTP_400_BAD_REQUEST

        raise HTTPException(
            status_code=status_code,
            detail=detail,
        )

    return AdminEnrollmentStudent(**student)


@router.delete(
    "/classrooms/{classroom_id}/students/{student_id}",
    response_model=AdminEnrollmentStudent,
)
async def remove_student_enrollment(
    classroom_id: int,
    student_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(
        require_roles(UserRole.ADMIN)
    ),
) -> AdminEnrollmentStudent:
    try:
        student = await remove_school_student_enrollment(
            db=db,
            school_id=current_user.school_id,
            classroom_id=classroom_id,
            student_id=student_id,
        )
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=str(exc),
        )

    return AdminEnrollmentStudent(**student)

